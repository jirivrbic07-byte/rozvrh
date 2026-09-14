"""Convert the supplied Excel timetable into compact JSON consumed by the web app."""

from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries


CLASSES = [
    "ZOP BS-3A",
    "III/1A/2026",
    "III/2A/2026",
    "III/3A/2026",
    "III/4A/2026",
    "TECC AB-26A",
    "SZZ-3A",
]

DAY_IDS = ["mon", "tue", "wed", "thu", "fri"]


def slug(value: str) -> str:
    value = value.lower().replace("/", "-")
    return re.sub(r"[^a-z0-9]+", "-", value).strip("-")


def parse_header_date(value: object) -> str:
    match = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", str(value or ""))
    if not match:
        raise ValueError(f"Neplatné datum v záhlaví: {value!r}")
    day, month, year = map(int, match.groups())
    return datetime(year, month, day).date().isoformat()


def parse_week_number(value: object) -> int:
    match = re.search(r"Týden\s*:\s*(\d+)", str(value or ""), re.IGNORECASE)
    if not match:
        raise ValueError(f"Neplatné číslo týdne v záhlaví: {value!r}")
    return int(match.group(1))


def main() -> None:
    if len(sys.argv) < 3:
        raise SystemExit(
            "Použití: extract_schedule.py INPUT.xlsx [DALŠÍ_INPUT.xlsx ...] OUTPUT.json"
        )

    sources = [Path(value) for value in sys.argv[1:-1]]
    output = Path(sys.argv[-1])
    weeks: list[dict[str, object]] = []
    seen_week_numbers: set[int] = set()

    for source in sources:
        workbook = openpyxl.load_workbook(source, data_only=False)
        for worksheet in workbook.worksheets:
            week_header = worksheet.cell(1, 1).value
            if not re.search(r"Týden\s*:", str(week_header or ""), re.IGNORECASE):
                continue
            week_number = parse_week_number(week_header)
            if week_number in seen_week_numbers:
                raise ValueError(f"Týden {week_number} je ve vstupních souborech dvakrát")
            seen_week_numbers.add(week_number)
            day_columns = [
                cell.column
                for cell in worksheet[2]
                if isinstance(cell.value, str)
                and re.match(r"^(Po|Út|St|Čt|Pá)\s", cell.value.strip())
            ][:5]
            if len(day_columns) != 5:
                raise ValueError(f"List {worksheet.title}: nenalezeno všech pět dnů")

            days = []
            for day_index, column in enumerate(day_columns):
                date_iso = parse_header_date(worksheet.cell(2, column).value)
                days.append({"id": DAY_IDS[day_index], "date": date_iso})

            merge_anchors: dict[tuple[int, int], tuple[int, int]] = {}
            for merged in worksheet.merged_cells.ranges:
                min_col, min_row, max_col, _ = range_boundaries(str(merged))
                merge_anchors[(min_row, min_col)] = (min_col, max_col)

            schedules: dict[str, list[dict[str, object]]] = {
                class_name: [] for class_name in CLASSES
            }
            class_rows = {
                str(cell.value).strip(): cell.row
                for cell in worksheet["A"]
                if str(cell.value).strip() in CLASSES
            }

            for class_name in CLASSES:
                row_number = class_rows.get(class_name)
                if not row_number:
                    continue

                block_index = 0
                for cell in worksheet[row_number][1:]:
                    if cell.value is None:
                        continue
                    min_col, max_col = merge_anchors.get(
                        (cell.row, cell.column), (cell.column, cell.column)
                    )
                    day_index = max(
                        (
                            i
                            for i, start_col in enumerate(day_columns)
                            if start_col <= min_col
                        ),
                        default=-1,
                    )
                    if day_index < 0:
                        continue
                    day_end = (
                        day_columns[day_index + 1] - 1
                        if day_index < 4
                        else worksheet.max_column
                    )
                    max_col = min(max_col, day_end)
                    slot_start = worksheet.cell(3, min_col).value
                    slot_end = worksheet.cell(3, max_col).value
                    if not isinstance(slot_start, int) or not isinstance(slot_end, int):
                        continue

                    block_index += 1
                    block_id = (
                        f"w{week_number}-{slug(class_name)}-{DAY_IDS[day_index]}-"
                        f"{slot_start}-{slot_end}-{block_index}"
                    )
                    schedules[class_name].append(
                        {
                            "id": block_id,
                            "dayId": DAY_IDS[day_index],
                            "slotStart": slot_start,
                            "slotEnd": slot_end,
                            "raw": re.sub(
                                r"\s+", " ", str(cell.value).replace("\u00a0", " ")
                            ).strip(),
                        }
                    )

            first_date = days[0]["date"]
            last_date = days[-1]["date"]
            weeks.append(
                {
                    "id": str(week_number),
                    "label": f"Týden {week_number}",
                    "startDate": first_date,
                    "endDate": last_date,
                    "days": days,
                    "schedules": schedules,
                }
            )

    payload = {
        "source": ", ".join(source.name for source in sources),
        "classes": CLASSES,
        "weeks": sorted(weeks, key=lambda week: int(str(week["id"]))),
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    json_text = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    output.write_text(json_text, encoding="utf-8")
    public_output = Path(__file__).resolve().parents[1] / "public" / "schedule.json"
    public_output.parent.mkdir(parents=True, exist_ok=True)
    public_output.write_text(json_text, encoding="utf-8")


if __name__ == "__main__":
    main()
